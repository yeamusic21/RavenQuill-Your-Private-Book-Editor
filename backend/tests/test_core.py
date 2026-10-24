from docx import Document

from ravenquill.docx_io import build_docx, parse_docx
from ravenquill.llm import strip_reasoning
from ravenquill.personas import check_edit, parse_batch
from ravenquill.textutil import diff_ops, from_segments, to_segments, typography

SETTINGS = {"book_font": "Garamond", "font_size": 11, "trim": "6x9", "scene_break": "* * *"}


def test_markup_roundtrip():
    text = "He was <i>very</i> late, <b><i>again</i></b>."
    assert from_segments(to_segments(text)) == text


def test_parse_docx(sample_docx):
    chapters = parse_docx(sample_docx)
    assert [c["title"] for c in chapters] == ["Chapter 1", "Chapter 2"]
    kinds = [p["kind"] for p in chapters[0]["paragraphs"]]
    assert kinds == ["heading", "body", "body", "scene_break", "body"]
    assert "<i>Faster</i>" in chapters[0]["paragraphs"][1]["text"]


def test_typography():
    assert typography('"Wait..." she said -- softly.') == "“Wait…” she said—softly."
    assert typography("It's the '90s and 'tis fine") == "It’s the ’90s and ’tis fine"
    assert typography("She said 'hello' to me.") == "She said ‘hello’ to me."


def test_strip_reasoning():
    assert strip_reasoning("<think>hmm</think>\nAnswer") == "Answer"
    assert strip_reasoning("hmm let me think</think>Answer") == "Answer"


def test_parse_batch():
    assert parse_batch("[1] One.\n\n[2] Two\ncontinued.", 2) == ["One.", "Two continued."]
    assert parse_batch("[1] One.", 2) is None


def test_guards():
    before = "He walked to the store and he bought some milk and then he walked home again."
    assert check_edit("prose", "standard", before, "He walked to the store, bought some milk, and walked home.") is None
    assert check_edit("prose", "standard", before, "A tapestry of errands unfolded.") is not None
    assert "formatting" in check_edit("copyedit", None, "A <i>big</i> dog ran.", "A big dog ran.")
    assert check_edit("prose", "standard", '"I aint going," he said.', '"I am not going," he said.') is not None


def test_diff_and_tracked_export(tmp_path):
    ops = diff_ops("the the dog", "the dog")
    assert any(o["op"] == "delete" for o in ops)
    pars = [
        {"kind": "heading", "original": "Chapter 1", "current": "Chapter 1"},
        {"kind": "body", "original": "Their was a <i>dog</i>.", "current": "There was a <i>dog</i>."},
        {"kind": "scene_break", "original": "***", "current": "***"},
        {"kind": "body", "original": "End.", "current": "End."},
    ]
    path = tmp_path / "t.docx"
    build_docx(path, "My Book", pars, SETTINGS, "tracked")
    xml = Document(path).element.xml
    assert "w:ins" in xml and "w:del" in xml
    build_docx(path, "My Book", pars, SETTINGS, "clean")
    texts = [p.text for p in Document(path).paragraphs]
    assert texts == ["My Book", "Chapter 1", "There was a dog.", "* * *", "End."]
