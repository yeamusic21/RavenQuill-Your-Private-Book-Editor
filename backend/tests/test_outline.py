import time

from docx import Document
from fastapi.testclient import TestClient

from ravenquill.personas import parse_outline_review

OUTLINE = """## Act 1
Mara finds a locked chest in her late mother's studio.

## Act 2
Mara leaves the city to find the locksmith.

## Act 3
Mara defeats the Duke with the sword."""


def test_parse_outline_review():
    reply = """CHECK:
[4] FIXED
[5] Still open, the outline doesn't say
**[7]** - fixed
CHECK: [8] FIXED
**HOLE:** Where does the sword come from?
WHERE: Act 3
SEVERITY: Major
WHY: The sword is never set up.
It appears from nowhere.

HOLE: Chest never opened
WHERE: Act 1
SEVERITY: minor
WHY: Set up, no payoff.
NO NEW HOLES"""
    fixed, holes = parse_outline_review(reply, [4, 5, 7, 8])
    assert fixed == [4, 7, 8]
    assert [h["title"] for h in holes] == ["Where does the sword come from?", "Chest never opened"]
    assert holes[0]["detail"] == "The sword is never set up. It appears from nowhere."
    assert holes[1]["detail"] == "Set up, no payoff."
    assert holes[0]["severity"] == "major" and holes[1]["severity"] == "minor"
    assert parse_outline_review("NO NEW HOLES", [1]) == ([], [])
    assert parse_outline_review("FIXED: 1, 9\nNO NEW HOLES", [1]) == ([1], [])  # 9 isn't an open hole


def wait_idle(client, outline_id, timeout=10):
    deadline = time.time() + timeout
    while time.time() < deadline:
        o = client.get(f"/api/outlines/{outline_id}").json()
        if o["status"] != "reviewing":
            return o
        time.sleep(0.1)
    raise AssertionError("review did not finish")


def test_outline_rounds(app_env, monkeypatch, tmp_path):
    prompts = []

    def fake_chat(self, system, user, max_tokens, temperature=None):
        from ravenquill.llm import Reply
        prompts.append(user)
        if len(prompts) == 1:
            text = ("HOLE: Sword from nowhere\nWHERE: Act 3\nSEVERITY: major\nWHY: Never set up.\n"
                    "HOLE: Why the locksmith?\nWHERE: Act 2\nSEVERITY: minor\nWHY: No reason given.")
        else:
            open_id = user.split("OPEN HOLES", 1)[1].split("[", 1)[1].split("]", 1)[0]
            # Re-raising a hole it just fixed (as small models do) must not create a duplicate.
            text = f"CHECK:\n[{open_id}] FIXED\nHOLE: Sword from nowhere!\nWHERE: Act 3\nWHY: Never set up."
        return Reply(text, 30, 0.1)

    monkeypatch.setattr("ravenquill.llm.LLM.chat", fake_chat)
    with TestClient(app_env.app) as client:
        o = client.post("/api/outlines", data={"name": "Quill", "text": OUTLINE}).json()
        oid = o["id"]
        assert o["words"] > 20 and o["holes"] == []

        client.post(f"/api/outlines/{oid}/review")
        o = wait_idle(client, oid)
        assert o["status"] == "idle", o
        assert len(o["holes"]) == 2 and all(h["status"] == "open" for h in o["holes"])
        sword, locksmith = o["holes"]  # majors first

        # Author explains one, fixes the outline for the other, and checks again.
        assert client.patch(f"/api/holes/{locksmith['id']}", json={"status": "explained"}).status_code == 400
        client.patch(f"/api/holes/{locksmith['id']}",
                     json={"status": "explained", "response": "Her mother's letter names him."})
        client.patch(f"/api/outlines/{oid}", json={"text": OUTLINE + "\nThe sword hangs in the studio in Act 1."})
        client.post(f"/api/outlines/{oid}/review")
        o = wait_idle(client, oid)

        assert "Her mother's letter names him." in prompts[1] and "SETTLED" in prompts[1]
        assert {h["id"]: h["status"] for h in o["holes"]} == {sword["id"]: "fixed", locksmith["id"]: "explained"}
        assert [r["new_holes"] for r in o["rounds"]] == [2, 0] and o["rounds"][1]["same_text"]

        report = client.get(f"/api/outlines/{oid}/export").text
        assert "## Fixed" in report and "Her mother's letter" in report

        # Upload as .docx too.
        doc = Document()
        doc.add_heading("Act 1", level=1)
        doc.add_paragraph("Mara finds a chest.")
        path = tmp_path / "outline.docx"
        doc.save(path)
        with open(path, "rb") as f:
            up = client.post("/api/outlines", data={"name": ""}, files={"file": ("My outline.docx", f)}).json()
        assert up["name"] == "My outline" and up["text"] == "## Act 1\n\nMara finds a chest."
        assert len(client.get("/api/outlines").json()) == 2
