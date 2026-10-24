import importlib
import re

import pytest
from docx import Document


@pytest.fixture
def sample_docx(tmp_path):
    """A tiny manuscript with chapters, italics, a scene break and some errors."""
    doc = Document()
    doc.add_heading("Chapter 1", level=1)
    p = doc.add_paragraph("Their was a dog named Rex and he run fast down the the road. ")
    p.add_run("Faster").italic = True
    p.add_run(" than anything.")
    doc.add_paragraph('"I aint going," said Mara. "Not today."')
    doc.add_paragraph("***")
    doc.add_paragraph("The next morning it was raining -- hard. She waited...")
    doc.add_heading("Chapter 2", level=1)
    doc.add_paragraph("Rex came back home at noon, tired and muddy.")
    path = tmp_path / "book.docx"
    doc.save(path)
    return path


@pytest.fixture
def app_env(tmp_path, monkeypatch):
    """Fresh data dir + a fake LLM so tests run instantly without LM Studio."""
    monkeypatch.setenv("RAVENQUILL_DATA", str(tmp_path / "data"))
    from ravenquill import config, db, pipeline, api
    for mod in (config, db, pipeline, api):
        importlib.reload(mod)

    def fake_chat(self, system, user, max_tokens, temperature=None):
        from ravenquill.llm import Reply
        if "EDIT:" in user:
            edit = user.split("EDIT:\n", 1)[1]
            # "Edit" by fixing a couple of known errors, keep the rest.
            fixed = edit.replace("Their was", "There was").replace("run fast", "ran fast").replace("the the", "the")
            return Reply(fixed, 50, 0.1)
        return Reply("- Plain first person past tense.\nCHARACTERS: Rex - dog", 20, 0.1)

    monkeypatch.setattr("ravenquill.llm.LLM.chat", fake_chat)
    db.init()
    return api
