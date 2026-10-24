"""Reading a manuscript .docx and writing book-formatted .docx files.

Exports open cleanly in Word and Google Docs (File > Open > Upload)."""
import re
from datetime import datetime, timezone

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from .textutil import SCENE_BREAK_RE, diff_ops, from_segments, to_segments, word_count

CHAPTER_RE = re.compile(
    r"^(chapter|prologue|epilogue|part|book|interlude|afterword|foreword|preface|introduction)\b"
    r"|^[IVXLC]+\.?$|^\d{1,3}\.?$",
    re.IGNORECASE,
)
SYNTHETIC_SECTION_WORDS = 2500


# --- import ----------------------------------------------------------------------

def _run_is(run, attr: str) -> bool:
    value = getattr(run, attr)
    if value is None and run.style is not None:
        value = getattr(run.style.font, attr)
    return bool(value)


def _paragraph_markup(paragraph) -> str:
    segments = [(run.text, _run_is(run, "italic"), _run_is(run, "bold")) for run in paragraph.runs]
    text = from_segments(segments)
    return re.sub(r"\s+", " ", text).strip()


def _is_heading(paragraph, text: str) -> bool:
    style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
    if style.startswith("heading") or style == "title":
        return True
    plain = re.sub(r"</?[ib]>", "", text)
    return len(plain) <= 60 and bool(CHAPTER_RE.match(plain.strip()))


def parse_docx(path) -> list[dict]:
    """Return chapters: [{'title': str, 'paragraphs': [{'kind', 'text'}]}]."""
    doc = Document(path)
    items = []
    for p in doc.paragraphs:
        text = _paragraph_markup(p)
        if not text:
            continue
        plain = re.sub(r"</?[ib]>", "", text)
        if SCENE_BREAK_RE.match(plain):
            items.append({"kind": "scene_break", "text": plain.strip()})
        elif _is_heading(p, text):
            items.append({"kind": "heading", "text": text})
        else:
            items.append({"kind": "body", "text": text})

    chapters: list[dict] = []
    if any(i["kind"] == "heading" for i in items):
        current = None
        for item in items:
            if item["kind"] == "heading":
                # Consecutive headings ("Chapter 1" + "The Storm") belong together.
                if current and all(x["kind"] == "heading" for x in current["paragraphs"]):
                    current["title"] += " — " + re.sub(r"</?[ib]>", "", item["text"])
                    current["paragraphs"].append(item)
                    continue
                current = {"title": re.sub(r"</?[ib]>", "", item["text"]), "paragraphs": [item]}
                chapters.append(current)
            else:
                if current is None:
                    current = {"title": "Opening", "paragraphs": []}
                    chapters.append(current)
                current["paragraphs"].append(item)
    else:
        # No headings found: split into sections so the editors have manageable units.
        current, words = None, 0
        for item in items:
            if current is None or words >= SYNTHETIC_SECTION_WORDS:
                current, words = {"title": f"Section {len(chapters) + 1}", "paragraphs": []}, 0
                chapters.append(current)
            current["paragraphs"].append(item)
            words += word_count(item["text"])
    return chapters


# --- export ----------------------------------------------------------------------

TRIMS = {
    "6x9": (Inches(6), Inches(9), Inches(0.75)),
    "5.5x8.5": (Inches(5.5), Inches(8.5), Inches(0.7)),
    "letter": (Inches(8.5), Inches(11), Inches(1)),
}


def _setup_styles(doc, settings):
    font_name, size = settings["book_font"], settings["font_size"]

    normal = doc.styles["Normal"]
    normal.font.name = font_name
    normal.font.size = Pt(size)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    pf = normal.paragraph_format
    pf.first_line_indent = Inches(0.3)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.15
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.widow_control = True

    first = doc.styles.add_style("First Paragraph", 1)  # 1 = paragraph style
    first.base_style = normal
    first.paragraph_format.first_line_indent = Inches(0)

    for name in ("Heading 1", "Title"):
        style = doc.styles[name]
        style.font.name = font_name
        style.font.bold = False
        style.font.italic = False
        style.font.color.rgb = None
        style.element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        hp = style.paragraph_format
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hp.first_line_indent = Inches(0)
        hp.keep_with_next = True
    doc.styles["Heading 1"].font.size = Pt(size + 7)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(72)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(30)
    doc.styles["Heading 1"].paragraph_format.page_break_before = True
    doc.styles["Title"].font.size = Pt(size + 17)

    scene = doc.styles.add_style("Scene Break", 1)
    scene.base_style = normal
    scene.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    scene.paragraph_format.first_line_indent = Inches(0)
    scene.paragraph_format.space_before = Pt(size)
    scene.paragraph_format.space_after = Pt(size)


def _add_runs(paragraph, text: str):
    for seg, italic, bold in to_segments(text):
        run = paragraph.add_run(seg)
        run.italic = italic or None
        run.bold = bold or None


def _tracked_run(tag: str, text: str, italic: bool, bold: bool, rev_id: int, date: str):
    """Build <w:ins> or <w:del> wrapping a single run."""
    wrapper = OxmlElement(f"w:{tag}")
    wrapper.set(qn("w:id"), str(rev_id))
    wrapper.set(qn("w:author"), "RavenQuill")
    wrapper.set(qn("w:date"), date)
    run = OxmlElement("w:r")
    if italic or bold:
        rpr = OxmlElement("w:rPr")
        if bold:
            rpr.append(OxmlElement("w:b"))
        if italic:
            rpr.append(OxmlElement("w:i"))
        run.append(rpr)
    t = OxmlElement("w:delText" if tag == "del" else "w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    wrapper.append(run)
    return wrapper


def _add_tracked_runs(paragraph, before: str, after: str, counter: list, date: str):
    for op in diff_ops(before, after):
        if op["op"] == "equal":
            run = paragraph.add_run(op["text"])
            run.italic = op["i"] or None
            run.bold = op["b"] or None
        else:
            counter[0] += 1
            tag = "ins" if op["op"] == "insert" else "del"
            paragraph._p.append(_tracked_run(tag, op["text"], op["i"], op["b"], counter[0], date))


def build_docx(path, title: str, paragraphs: list[dict], settings: dict, mode: str = "clean"):
    """paragraphs: dicts with kind/original/current. mode: 'clean' | 'tracked' | 'original'."""
    doc = Document()
    _setup_styles(doc, settings)
    width, height, margin = TRIMS.get(settings.get("trim", "6x9"), TRIMS["6x9"])
    for section in doc.sections:
        section.page_width, section.page_height = width, height
        section.left_margin = section.right_margin = margin
        section.top_margin = section.bottom_margin = margin

    # Title page
    title_par = doc.add_paragraph(title, style="Title")
    title_par.paragraph_format.space_before = Pt(160)
    if paragraphs and paragraphs[0]["kind"] != "heading":
        title_par.add_run().add_break(WD_BREAK.PAGE)

    counter, date = [0], datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    previous_kind = "heading"
    for i, par in enumerate(paragraphs):
        kind = par["kind"]
        text = par["original"] if mode == "original" else par["current"]
        if kind == "heading":
            # Only the first of consecutive headings breaks the page.
            style = "Heading 1"
            p = doc.add_paragraph(style=style)
            if previous_kind == "heading" and i > 0:
                p.paragraph_format.page_break_before = False
                p.paragraph_format.space_before = Pt(0)
        elif kind == "scene_break":
            p = doc.add_paragraph(settings.get("scene_break") or text, style="Scene Break")
            previous_kind = kind
            continue
        else:
            p = doc.add_paragraph(style="First Paragraph" if previous_kind in ("heading", "scene_break") else None)

        if mode == "tracked" and par["original"] != par["current"]:
            _add_tracked_runs(p, par["original"], par["current"], counter, date)
        else:
            _add_runs(p, text)
        previous_kind = kind

    doc.save(path)
    return path
